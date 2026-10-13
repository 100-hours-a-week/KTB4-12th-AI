"""프로파일링 코어 30개 상품 콜백 검증 — 시나리오 엑셀(가상 수신자) → 앱 전체 → 결과 화면 + 파일.

    cd profiling && uv run python -m tools.scenario_callback30.run [--modes exclude,penalty] [--keep] [--top 5]

한 사람마다 실제 앱 경로를 그대로 탄다: 7.6 POST(TestClient) → 202 → Supervisor 워커 → pipeline.profile()(별점 편향 · 풀 30개)
→ PostgreSQL 저장(프로필 + 리뷰) → **진짜 HttpBackendPort 가 보내는 7.7 HTTP 요청**(MockTransport 가 받아 경로·토큰·본문을 기록).
서버를 따로 띄우지 않고, LLM 도 부르지 않는다. 카탈로그는 로컬 DB(`ai_catalog`, BE 번호)를 쓴다.

결과는 세 가지로 본다.
  화면  사람마다: 시간(202 · 콜백 · 서버) · 소분류 점수 · 태그 · 상위 N 과 점수 내역 · 검사 결과
  엑셀  tools/scenario_callback30/out/시나리오_결과_<시각>.xlsx — 요약 · 점수_책정 · 추천30(점수 내역) · 비선호_위치 · 태그 · 콜백_원문 · 검사 · 입력
  JSON  같은 이름 .json
검사는 두 갈래다. **자동 검사**(모든 실행): 콜백 수신 · 본문 계약 · 개수 · 카탈로그·재고 · 재현성 · 비선호 · 점수 공식 독립 재계산 ·
점수 저장 · 리뷰 행 · 원문 미저장 · 실행 기록. **시나리오 기대**: 엑셀 '기대' 시트. 하나라도 실패하면 종료 코드 1.
"""

from __future__ import annotations

import argparse
import json
import logging
import os
import re
import sys
import time
import warnings
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

import httpx
import sqlalchemy as sa

warnings.filterwarnings("ignore", message="Using `httpx` with `starlette.testclient`")   # 도구 출력에 섞이는 의존성 경고

HERE = Path(__file__).resolve().parent
DEFAULT_SCENARIOS = HERE / "scenarios.xlsx"
DEFAULT_OUT = HERE / "out"
PATH_76 = "/api/internal/v1/ai/profile/extract-and-pool"
TOKEN = "scenario-token"
RID_BASE = {"exclude": 9_300_000, "penalty": 9_310_000}
MODES = ("exclude", "penalty")


# ---------------------------------------------------------------------------
# 시나리오 읽기
# ---------------------------------------------------------------------------


@dataclass
class Expect:
    mode: str
    kind: str
    target: str
    value: str
    note: str
    cond: str = ""          # '태그=끔' · '태그=켬' · '고르기=d' · '모델=gemma' (';' 로 여럿) — 비면 언제나
    grade: str = "필수"      # 필수(실패하면 하네스 실패) · 관찰(보고만)
    nature: str = ""        # 계약 · 별점 · 회귀 · 태그품질 · 취향 · 섞임


OBSERVE = "기대·관찰"


def blocking(check: tuple) -> bool:
    """관찰 등급 기대는 실패해도 하네스를 실패시키지 않는다 — 모델 · 프롬프트 품질 비교용."""
    return check[0] != OBSERVE


def applies(e: Expect, config: dict[str, str]) -> bool:
    """기대의 조건이 이 실행 설정(태그 · 고르기 · 모델)에 맞나."""
    for part in (x.strip() for x in e.cond.split(";") if x.strip()):
        key, _, want = part.partition("=")
        have = {"태그": "켬" if config.get("tagger", "none") != "none" else "끔", "고르기": config.get("strategy", "d"),
                "모델": config.get("model", "")}.get(key.strip(), "")
        if (key.strip() == "모델" and want.strip() not in have) or (key.strip() != "모델" and have != want.strip()):
            return False
    return True


@dataclass
class User:
    no: str
    alias: str
    persona: str
    disliked: list[tuple[int, str]]
    preference: str | None
    note: str
    reviews: list[dict[str, Any]] = field(default_factory=list)
    expects: list[Expect] = field(default_factory=list)

    def body(self, rid: int, sv: int = 1) -> dict[str, Any]:
        return {"recipientUserId": rid, "sourceVersion": sv,
                "dislikedCategories": [{"categoryId": i, "categoryName": n} for i, n in self.disliked],
                "giftPreference": self.preference, "reviews": self.reviews}


def load_scenarios(path: Path) -> list[User]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    rows = lambda name: [r for r in wb[name].iter_rows(min_row=2, values_only=True) if r and r[0]]
    users: dict[str, User] = {}
    for no, alias, persona, disliked, pref, note in rows("사용자"):
        pairs = []
        for part in str(disliked or "").split(";"):
            if part.strip():
                cid, _, name = part.partition(":")
                pairs.append((int(cid), name.strip()))
        users[no] = User(no, alias or "", persona or "", pairs, (pref or None), note or "")
    for no, pos, pid, rating, text in sorted(rows("리뷰"), key=lambda r: (r[0], r[1])):
        users[no].reviews.append({"productId": int(pid), "rating": int(rating), "reviewText": text or None})
    for row in rows("기대"):
        no, mode, kind, target, value, note, *rest = list(row) + [None] * 3
        cond, grade, nature = (rest + [None, None, None])[:3]
        users[no].expects.append(Expect(mode, kind, str(target or ""), str(value or ""), note or "", str(cond or ""), str(grade or "필수"),
                                        str(nature or "")))
    return list(users.values())


# ---------------------------------------------------------------------------
# 한 모드 실행 — 앱 전체를 TestClient 로
# ---------------------------------------------------------------------------


class _Collect(logging.Handler):
    """profiling.* 경고를 수신자 번호별로 모은다 (화면·엑셀에 '서버 경고'로)."""

    def __init__(self) -> None:
        super().__init__(logging.WARNING)
        self.records: list[str] = []

    def emit(self, record: logging.LogRecord) -> None:
        self.records.append(record.getMessage())


def run_mode(mode: str, users: list[User], engine: sa.Engine, model_env: dict[str, str] | None = None) -> list[dict[str, Any]]:
    """model_env = 태그 추출 · 임베딩 설정(PROFILING_TAGGER · TAGGER_MODEL · EMBEDDING_DIR). 없으면 꺼진 채(태그 없이 D)."""
    os.environ.update({"PROFILING_CATALOG_SOURCE": "db", "PROFILING_SERVICE_TOKEN": TOKEN,
                       "PROFILING_DISLIKE_MODE": mode, "PROFILING_LOG_LEVEL": "ERROR", "PROFILING_TAGGER": "none"})
    os.environ.pop("PROFILING_EMBEDDING_DIR", None)
    os.environ.update(model_env or {})
    from fastapi.testclient import TestClient

    from profiling import pipeline
    from profiling.backend import HttpBackendPort
    from profiling.settings import get_settings
    from profiling.types import DislikedCategory, ProfileRequest, Review

    get_settings.cache_clear()
    from profiling.main import app

    calls: list[dict[str, Any]] = []

    def backend_handler(request: httpx.Request) -> httpx.Response:     # 7.7 을 받는 쪽(BE 흉내) — 실제 HTTP 요청을 그대로 기록
        body = json.loads(request.content)
        calls.append({"t": time.perf_counter(), "path": request.url.path, "auth": request.headers.get("authorization"), "body": body})
        return httpx.Response(200, json={"message": "ok", "data": {"recipientUserId": body["recipientUserId"],
                                                                  "sourceVersion": body["sourceVersion"], "profileStatus": "COMPLETED"}})

    results = []
    plog = logging.getLogger("profiling")
    collector = _Collect()
    plog.addHandler(collector)
    propagate = plog.propagate
    plog.propagate = False                              # 앱 로그(시작·종료 포함)는 화면에 섞지 않고 '서버 경고'로만 모은다
    with TestClient(app) as client:
        app.state.backend = HttpBackendPort("http://backend.scenario", TOKEN, 5.0, transport=httpx.MockTransport(backend_handler))
        spy = _SpyTagger(app.state.tagger) if app.state.tagger is not None else None
        app.state.tagger = spy                          # LLM 원출력을 기록 — 하네스가 같은 검증기로 태그를 다시 만든다(DB 엔 축 · 프롬프트판이 아직 없음, 0006)
        plog.setLevel(logging.INFO)
        catalog = app.state.catalog
        cv, products = catalog.active()
        strategy = pipeline.PoolStrategy(get_settings().POOL_STRATEGY)
        for i, u in enumerate(users, start=1):
            rid = RID_BASE[mode] + i
            body = u.body(rid)
            collector.records.clear()
            n_calls = len(calls)
            t0 = time.perf_counter()
            resp = client.post(PATH_76, json=body, headers={"Authorization": f"Bearer {TOKEN}"})
            t202 = time.perf_counter() - t0
            wait_s = 15 + (get_settings().TAGGER_TIMEOUT_S if get_settings().TAGGER != "none" else 0)   # LLM 이 켜지면 그만큼 더 기다린다
            deadline = t0 + wait_s
            while len(calls) == n_calls and time.perf_counter() < deadline:
                time.sleep(0.001)
            mine = [c for c in calls[n_calls:] if c["body"].get("recipientUserId") == rid]
            t_cb = (mine[0]["t"] - t0) if mine else None
            run_row = _wait_run(engine, rid, timeout=wait_s)
            app_warnings = list(collector.records)      # 앱이 낸 경고만 — 아래 하네스 계산이 같은 경고를 또 내므로 먼저 떠 둔다
            # 같은 입력으로 코드가 낼 순위(전체) — 콜백과 대조하고 점수 내역을 보여 주는 데 쓴다
            rq = ProfileRequest(recipient_user_id=rid, source_version=1, gift_preference=u.preference,
                                disliked_categories=[DislikedCategory(category_id=c, category_name=n) for c, n in u.disliked],
                                reviews=[Review(product_id=r["productId"], rating=r["rating"], review_text=r["reviewText"]) for r in u.reviews])
            scores = pipeline.category_scores(rq, products, catalog)
            ranked = pipeline.rank_products(rq, products, dislike_mode=pipeline.DislikeMode(mode), scores=scores)
            # 같은 입력으로 코드가 고를 30개(D: 비선호 하드 제외 · 한 소분류 최대 10 · 같은 모델 하나) — 콜백과 대조하고 점수 내역을 보여 준다
            # 태그 — 서비스가 받은 LLM 원출력(spy)을 같은 검증기로(LLM 을 다시 부르지 않는다). 저장된 프로필 태그와도 대조한다(검사)
            prof = app.state.recipient_store.get(rid)
            tags, prompt_version = _tags_from_spy(spy, rq, products)
            usage = _usage_from_spy(spy, rq, products)
            llm_prompt = _prompt_sent(spy, rq, products)          # LLM 에 실제로 보낸 프롬프트(없으면 None — 안 불렀음)
            search, rule = pipeline.make_pool(rq, products, 30, cv, dislike_mode=pipeline.DislikeMode(mode), scores=scores,
                                              pool_strategy=strategy, tags=tags, embedding=app.state.embedding)
            at = {s.product.productId: (k, s) for k, s in enumerate(ranked, start=1)}
            pool = [at[i] for i in search.product_ids]           # (전체 순위, ScoredProduct) — 고른 순서
            verify = _verify_d(mine, rq, products, cv, mode, scores, strategy, tags, app.state.embedding, ranked) if strategy is pipeline.PoolStrategy.D else []
            results.append({"user": u, "mode": mode, "rid": rid, "status_202": resp.status_code, "t202": t202, "t_cb": t_cb,
                            "calls": mine, "run": run_row, "scores": scores, "ranked": ranked, "products": products,
                            "pool": pool, "rule": rule, "strategy": strategy, "tags": tags, "verify": verify,
                            "config": {"tagger": get_settings().TAGGER, "model": get_settings().TAGGER_MODEL if get_settings().TAGGER != "none" else "",
                                       "strategy": strategy.value},
                            "prompt_version": prompt_version, "stored_tags": list(prof.preferred_tags) if prof else None,
                            "usage": usage, "llm_prompt": llm_prompt, "warnings": app_warnings})
    plog.removeHandler(collector)
    plog.propagate = propagate
    return results


class _SpyTagger:
    """ports.TagExtractor 를 감싸 원출력을 프롬프트별로 기록한다(시험용). 동작은 그대로 넘긴다."""

    def __init__(self, inner) -> None:
        self.inner, self.model, self.raws, self.usages = inner, inner.model, {}, {}

    def extract(self, system: str, prompt: str, schema: dict):
        reply = self.inner.extract(system, prompt, schema)
        self.raws[prompt] = getattr(reply, "text", reply)
        self.usages[prompt] = getattr(reply, "usage", None)   # 토큰 · 시간(로컬도 잰다 — 외부 API 비용 추정용)
        return reply

    def close(self) -> None:
        close = getattr(self.inner, "close", None)
        if callable(close):
            close()


def _verify_d(calls, rq, products, cv, mode, scores, strategy, tags, embedding, ranked) -> list[tuple[str, bool, str]]:
    """D 독립 판정(verify_d — selection 을 부르지 않음) + 입력을 바꿔 보는 검사 두 가지. 대상은 **서비스가 보낸 콜백 30개**."""
    from profiling import pipeline
    from tools.scenario_callback30 import verify_d

    ids = calls[0]["body"]["recommendedProductIds"] if calls else []
    dm = pipeline.DislikeMode(mode)
    out = verify_d.verify(ids, ranked, tags, embedding, set(pipeline.reviewed_ids(rq)), dm is pipeline.DislikeMode.PENALTY)
    run = lambda q: pipeline.make_pool(q, products, 30, cv, dislike_mode=dm, scores=scores, pool_strategy=strategy, tags=tags,
                                       embedding=embedding)[0].product_ids
    a, b = run(rq), run(rq)
    out.append(("입력 변형 · 결정성: 같은 입력 두 번 → 같은 30개", a == b, ""))
    by_id = {p.productId: p for p in products}
    top = next((by_id[i] for i in ids if i in by_id and by_id[i].parentCategoryId), None)
    if top is not None and len(rq.disliked_categories) < 5:
        from profiling.types import DislikedCategory
        rq2 = rq.model_copy(update={"disliked_categories": [*rq.disliked_categories,
                                                            DislikedCategory(category_id=top.parentCategoryId, category_name=top.parentCategoryName)]})
        left = [i for i in run(rq2) if by_id[i].parentCategoryId == top.parentCategoryId]
        out.append((f"입력 변형 · 비선호 추가({top.parentCategoryName}) → 그 대분류 상품이 사라짐", not left, f"남은 {len(left)}개"))
    return out


def _tags_from_spy(spy, rq, products) -> tuple[list, str | None]:
    """서비스가 그 요청에 받은 LLM 원출력 → 같은 검증기 → 풀 고르기 태그 · prompt_version. 추출기가 꺼졌거나 호출이 없었으면 ([], None)."""
    from profiling import tagging

    if spy is None:
        return [], None
    prep = tagging.prepare(rq, {p.productId: p for p in products})
    prompt, ratings = prep.prompt, prep.ratings
    raw = spy.raws.get(prompt)
    if raw is None:
        return [], None
    keep, drop = tagging.validate(tagging.parse(raw), ratings, prep.has_pref)
    return tagging.tags_of(tagging.to_validation(keep, drop, [])), f"{tagging.PROMPT_VERSION}@{spy.model}"


def _usage_from_spy(spy, rq, products):
    """그 요청의 LLM 사용량(ports.TokenUsage) — 추출기가 꺼졌거나 호출이 없었으면 None."""
    from profiling import tagging

    if spy is None:
        return None
    return spy.usages.get(tagging.build_prompt(rq, {p.productId: p for p in products})[0])


def _prompt_sent(spy, rq, products) -> str | None:
    """그 요청에 LLM 으로 보낸 프롬프트 — 가드레일 정리 뒤의 글. 추출기가 꺼졌거나 부르지 않았으면 None."""
    from profiling import tagging

    if spy is None:
        return None
    prompt = tagging.prepare(rq, {p.productId: p for p in products}).prompt
    return prompt if prompt in spy.raws else None


def usage_summary(results) -> dict[str, Any] | None:
    """LLM 사용량 요약 — 호출 수 · 입력/출력 토큰 합 · 사람당 p50/p95 · 시간 p95. 로컬이라 돈은 0, 외부 API 로 바꿀 때 단가만 곱하면 된다."""
    us = [r["usage"] for r in results if r.get("usage") is not None]
    if not us:
        return None

    def pct(xs, q):
        xs = sorted(x for x in xs if x is not None)
        return xs[min(len(xs) - 1, round(q * (len(xs) - 1)))] if xs else None

    ins, outs = [u.input_tokens for u in us], [u.output_tokens for u in us]
    return {"calls": len(us), "input_tokens": sum(x or 0 for x in ins), "output_tokens": sum(x or 0 for x in outs),
            "input_p50": pct(ins, .5), "input_p95": pct(ins, .95), "input_max": pct(ins, 1),
            "output_p50": pct(outs, .5), "output_p95": pct(outs, .95), "output_max": pct(outs, 1),
            "llm_ms_p50": pct([u.total_ms for u in us], .5), "llm_ms_p95": pct([u.total_ms for u in us], .95),
            "load_ms_max": pct([u.load_ms for u in us], 1), "stops": dict(Counter(u.stop for u in us)),
            "cost_usd": round(sum(u.cost_usd for u in us if u.cost_usd is not None), 6) if any(u.cost_usd is not None for u in us) else None}


def _wait_run(engine: sa.Engine, rid: int, timeout: float = 15.0) -> dict[str, Any] | None:
    t0 = time.perf_counter()
    while time.perf_counter() - t0 < timeout:
        with engine.connect() as c:
            row = c.execute(sa.text("""select status, callback_attempts, callback_payload->'recommendedProductIds' as ids,
                                       extract(epoch from updated_at - created_at) as srv from ai_profile.profile_runs
                                       where recipient_user_id=:r and source_version=1"""), {"r": rid}).mappings().first()
        if row and row["status"] in ("DELIVERED", "SUPERSEDED", "FAILED"):
            return dict(row)
        time.sleep(0.002)
    return None


# ---------------------------------------------------------------------------
# 검사
# ---------------------------------------------------------------------------


def _target_ids(spec: str) -> tuple[str, set[int]]:
    level, _, ids = spec.partition(":")
    return level.strip(), {int(x) for x in ids.split(",") if x.strip()}


def _match(p, level: str, ids: set[int]) -> bool:
    return (p.categoryId in ids) if level == "소분류" else (p.parentCategoryId in ids)


def check_run(r: dict[str, Any], engine: sa.Engine, other: dict[str, Any] | None) -> list[tuple[str, str, bool, str]]:
    """(종류, 검사, 통과, 내용) 목록."""
    from pydantic import ValidationError

    from profiling import pipeline
    from profiling.backend import CALLBACK_PATH
    from profiling.pipeline import BIAS_WEIGHT, DISLIKED_CATEGORY_SCORE
    from profiling.schemas import ProfileCallbackRequest
    from profiling.stores import DbRecipientProfileStore
    from profiling.types import Review, review_hash

    u, mode, rid = r["user"], r["mode"], r["rid"]
    by_id = {p.productId: p for p in r["products"]}
    out: list[tuple[str, str, bool, str]] = []
    add = lambda name, ok, detail="": out.append(("자동", name, bool(ok), detail))

    add("7.6 접수 202", r["status_202"] == 202, f"HTTP {r['status_202']}")
    calls = r["calls"]
    ok = len(calls) == 1 and calls[0]["path"] == CALLBACK_PATH.format(recipientUserId=rid) and calls[0]["auth"] == f"Bearer {TOKEN}"
    add("7.7 콜백 1회 · 경로 · 서비스 토큰", ok, f"{len(calls)}회 · {calls[0]['path'] if calls else '-'}")
    ids: list[int] = calls[0]["body"].get("recommendedProductIds", []) if calls else []
    try:
        cb = ProfileCallbackRequest.model_validate(calls[0]["body"]) if calls else None
        add("7.7 본문 계약(모르는 필드 없음 · 번호 일치 · 태그 없음)", cb is not None and cb.recipientUserId == rid and cb.sourceVersion == 1,
            json.dumps(calls[0]["body"], ensure_ascii=False)[:120] if calls else "")
    except ValidationError as e:
        add("7.7 본문 계약(모르는 필드 없음 · 번호 일치 · 태그 없음)", False, str(e)[:200])
    want_n = min(30, len(r["ranked"]))
    add("추천 개수 = min(30, 후보 수) · 중복 없음", len(ids) == want_n and len(set(ids)) == len(ids), f"{len(ids)}개(기대 {want_n})")
    reviewed = {rv["productId"] for rv in u.reviews}
    add("리뷰한 상품(이미 받은 선물)은 추천하지 않음", not (set(ids) & reviewed), f"겹침 {sorted(set(ids) & reviewed) or '-'}")
    bad = [i for i in ids if i not in by_id or by_id[i].availability == "unavailable"]
    add("전부 카탈로그에 있고 재고 없음(unavailable) 아님", not bad, f"문제 {bad[:5]}")
    run_ids = list((r["run"] or {}).get("ids") or [])
    top = [s.product.productId for _, s in r["pool"]]
    add(f"재현성: 콜백 30 = 실행 기록 30 = 같은 입력으로 코드가 고른 30({r['rule']})", ids == run_ids == top,
        "" if ids == top else f"첫 차이 {next((k for k, (a, b) in enumerate(zip(ids, top)) if a != b), '-')}위")

    dis_ids = {c for c, _ in u.disliked}
    dis_names = {n for _, n in u.disliked}
    is_dis = lambda p: (p.categoryId in dis_ids or p.categoryName in dis_names or p.parentCategoryId in dis_ids or p.parentCategoryName in dis_names)
    flags = [is_dis(by_id[i]) for i in ids if i in by_id]
    n_ok = sum(not s.disliked for s in r["ranked"])
    if mode == "exclude":
        add("비선호(제외 모드): 30개에 비선호 상품 0", not any(flags), f"비선호 {sum(flags)}개")
    else:
        first_dis = next((k for k, f in enumerate(flags) if f), None)
        add("비선호(감점 모드): 비선호 상품은 비선호 아닌 상품보다 뒤에만", first_dis is None or not any(not f for f in flags[first_dis:]),
            f"30개 중 비선호 {sum(flags)}개")
        if r["strategy"] is pipeline.PoolStrategy.D:
            add("비선호(D): 비선호 아닌 상품으로 채울 수 있으면 0 — 모자랄 때만 채움", sum(flags) == max(0, want_n - n_ok),
                f"비선호 {sum(flags)}개 · 비선호 아닌 후보 {n_ok}")
    for name, ok, detail in r.get("verify", []):
        add(name, ok, detail)
    if r["prompt_version"]:
        add("저장된 선호 태그 = 같은 원출력을 검증기로 다시 만든 태그", r["stored_tags"] == [t.text for t in r["tags"]],
            f"저장 {r['stored_tags']} · 재생성 {[t.text for t in r['tags']]}")
    if r["strategy"] is pipeline.PoolStrategy.D:
        from collections import Counter

        from profiling.selection import CAT_CAP, dup_key
        cnt = Counter(by_id[i].categoryId for i in ids if i in by_id)
        room = sum(min(CAT_CAP, v) for v in Counter(s.product.categoryId for s in r["ranked"] if not s.disliked).values())
        add(f"한 소분류 최대 {CAT_CAP}(D)", max(cnt.values(), default=0) <= CAT_CAP or room < want_n,
            f"최다 {max(cnt.values(), default=0)} · 소분류 {len(cnt)}종")
        keys = [dup_key(by_id[i].name, by_id[i].brand) for i in ids if i in by_id]
        keys = [k for k in keys if k is not None]
        add("같은 브랜드 + 모델 코드 중복 없음(D)", len(keys) == len(set(keys)), f"모델 코드 있는 상품 {len(keys)}개")

    # 점수 독립 재계산 — 코드(score_categories · rank_products)와 별개로 리뷰·카탈로그에서 직접
    seen, manual = set(), {}
    for rv in u.reviews:
        if rv["productId"] in seen:
            continue
        seen.add(rv["productId"])
        p = by_id.get(rv["productId"])
        if p is None:
            continue
        n, s = manual.get(p.categoryId, (0, 0))
        manual[p.categoryId] = (n + 1, s + rv["rating"] - 3)
    got = {s.category_id: (s.n, s.score_sum) for s in r["scores"]}
    add("소분류 점수(n · 점수 합) 독립 재계산과 같음", got == manual, f"코드 {got} · 직접 {manual}")
    cand = [p for p in r["products"] if p.availability != "unavailable"]
    # 인기 = 조회수 순위 백분위 × 2 — 조회수별 개수를 세어 "나보다 낮은 후보 수"를 직접 만든다(코드의 bisect 와 다른 방법)
    counts: dict[int, int] = {}
    for p in cand:
        counts[p.viewCount] = counts.get(p.viewCount, 0) + 1
    lower, acc = {}, 0
    for v in sorted(counts):
        lower[v], acc = acc, acc + counts[v]
    span = len(cand) - 1
    worst = 0.0
    for s in r["ranked"]:
        n, sm = manual.get(s.product.categoryId, (0, 0))
        expect = (2.0 * lower[s.product.viewCount] / span if span > 0 else 0.0) + BIAS_WEIGHT * (sm / (n + 2) if n else 0.0) + (DISLIKED_CATEGORY_SCORE if mode == "penalty" and is_dis(s.product) else 0.0)
        worst = max(worst, abs(expect - s.total))
    ordered = all((a.total, -a.product.productId) >= (b.total, -b.product.productId) for a, b in zip(r["ranked"], r["ranked"][1:]))
    add("점수 공식 독립 재계산(인기 백분위 + 편향 + 감점) · 정렬", worst < 1e-9 and ordered, f"최대 오차 {worst:.2e}")

    prof = DbRecipientProfileStore(engine).get(rid)
    add("저장된 category_scores = 코드 점수", prof is not None and prof.category_scores == r["scores"], "")
    seen, want_rows = set(), []
    for pos, rv in enumerate(u.reviews):
        if rv["productId"] in seen:
            continue
        seen.add(rv["productId"])
        want_rows.append((rv["productId"], rv["rating"], pos, review_hash(Review(product_id=rv["productId"], rating=rv["rating"], review_text=rv["reviewText"]))))
    have_rows = [(x.product_id, x.rating, x.position, x.review_hash) for x in (prof.reviews if prof else [])]
    add("리뷰 행(중복 제거 · 자리 · 지문) 저장", have_rows == want_rows, f"{len(have_rows)}행(기대 {len(want_rows)})")
    texts = [rv["reviewText"] for rv in u.reviews if rv["reviewText"]] + ([u.preference] if u.preference else [])
    with engine.connect() as c:
        dump = c.execute(sa.text(
            "select coalesce((select row_to_json(p)::text from ai_profile.recipient_profiles p where p.recipient_user_id=:r),'')"
            " || coalesce((select json_agg(x)::text from ai_profile.recipient_reviews x where x.recipient_user_id=:r),'')"
            " || coalesce((select json_agg(x)::text from ai_profile.profile_runs x where x.recipient_user_id=:r),'')"), {"r": rid}).scalar()
    leaked = [t for t in texts if t in dump]
    add("원문 미저장(리뷰 글 · 취향 문장이 DB 어디에도 없음)", not leaked, f"검사 문자열 {len(texts)}개 · 발견 {len(leaked)}")
    run = r["run"] or {}
    add("실행 기록 DELIVERED · 콜백 시도 1", run.get("status") == "DELIVERED" and run.get("callback_attempts") == 1, f"{run.get('status')} · {run.get('callback_attempts')}")

    # 시나리오 기대
    texts = [t.text for t in r["tags"]]
    for e in u.expects:
        if e.mode not in ("both", mode) or not applies(e, r["config"]):
            continue
        name = f"[{e.nature or '-'}] {e.kind} {e.target} {e.value} — {e.note}".strip()
        label = OBSERVE if e.grade == "관찰" else "기대·필수"
        prods = [by_id[i] for i in ids if i in by_id]
        n0 = len(out)
        if e.kind in ("상위N_최소", "상위N_최대"):
            level, tids = _target_ids(e.target)
            topn, k = (int(x) for x in e.value.split(":"))
            cnt = sum(1 for p in prods[:topn] if _match(p, level, tids))
            ok = cnt >= k if e.kind == "상위N_최소" else cnt <= k
            out.append(("기대", name, ok, f"상위 {topn} 중 {cnt}개"))
        elif e.kind == "조회수순":
            views = [p.viewCount for p in prods]
            out.append(("기대", name, views == sorted(views, reverse=True), f"조회수 {views[:5]}…"))
        elif e.kind == "모드간_동일":
            other_ids = other["calls"][0]["body"]["recommendedProductIds"] if other and other["calls"] else None
            out.append(("기대", name, other_ids is not None and other_ids == ids, "다른 모드 실행 없음" if other_ids is None else ""))
        elif e.kind == "순위_앞섬":
            a, _, b = e.target.partition(">")
            rank = {}
            for spec in (a, b):
                level, tids = _target_ids(spec)
                rank[spec] = next((k + 1 for k, s in enumerate(r["ranked"]) if _match(s.product, level, tids)), None)
            ok = rank[a] is not None and (rank[b] is None or rank[a] < rank[b])
            out.append(("기대", name, ok, f"{a} {rank[a]}위 · {b} {rank[b]}위 / {len(r['ranked'])}"))
        elif e.kind in ("태그_포함", "태그_제외"):
            hit = [t for t in texts if re.search(e.target, t)]
            ok = bool(hit) if e.kind == "태그_포함" else not hit
            out.append(("기대", name, ok, f"태그 {texts or '[]'} · 맞은 것 {hit or '-'}"))
        elif e.kind == "태그_없음":
            out.append(("기대", name, not texts, f"태그 {texts or '[]'}"))
        elif e.kind == "LLM_호출_없음":                      # 가드레일: 근거가 될 것이 없으면 LLM 을 부르지 않는다(10-13)
            out.append(("기대", name, r.get("llm_prompt") is None, "호출 없음" if r.get("llm_prompt") is None else "호출됨"))
        elif e.kind == "프롬프트_제외":                       # 가드레일: LLM 으로 보낸 글에 대상(정규식)이 없어야 한다 — 개인정보 · 욕설
            sent = r.get("llm_prompt") or ""
            hit = re.findall(e.target, sent)
            out.append(("기대", name, not hit, f"보낸 프롬프트 {len(sent)}자 · 발견 {len(hit)}" + ("" if sent else " (호출 없음)")))
        elif e.kind == "이름_최소":
            topn, k = (int(x) for x in e.value.split(":"))
            cnt = sum(1 for p in prods[:topn] if re.search(e.target, p.name))
            out.append(("기대", name, cnt >= k, f"상위 {topn} 중 {cnt}개"))
        elif e.kind == "종수_최소":
            topn, k = (int(x) for x in e.value.split(":"))
            cnt = len({p.categoryId for p in prods[:topn]})
            out.append(("기대", name, cnt >= k, f"상위 {topn} 의 소분류 {cnt}종"))
        else:
            out.append(("기대", name, False, f"모르는 검사 종류 '{e.kind}'"))
        out[n0:] = [(label, *c[1:]) for c in out[n0:]]          # 등급을 종류에 싣는다 — 필수만 하네스 실패로 센다
    return out


# ---------------------------------------------------------------------------
# 화면 · 파일
# ---------------------------------------------------------------------------


def ms(x: float | None) -> str:
    return "-" if x is None else f"{x * 1000:.1f}ms"


def print_run(r: dict[str, Any], checks, top: int) -> None:
    u = r["user"]
    auto = [c for c in checks if c[0] == "자동"]
    exp = [c for c in checks if c[0] == "기대·필수"]
    obs = [c for c in checks if c[0] == OBSERVE]
    a_ok, e_ok, o_ok = sum(c[2] for c in auto), sum(c[2] for c in exp), sum(c[2] for c in obs)
    mark = "✅" if a_ok == len(auto) and e_ok == len(exp) else "❌"
    ids = r["calls"][0]["body"]["recommendedProductIds"] if r["calls"] else []
    print(f"\n{mark} [{u.no} {u.alias} · {r['mode']}] 202 {ms(r['t202'])} · 7.7 콜백까지 {ms(r['t_cb'])} · 서버 {ms((r['run'] or {}).get('srv') and float(r['run']['srv']))}"
          f" · 상품 {len(ids)}개 · 자동 {a_ok}/{len(auto)} · 기대(필수) {e_ok}/{len(exp)}" + (f" · 관찰 {o_ok}/{len(obs)}" if obs else ""))
    if r["scores"]:
        parts = []
        for s in r["scores"]:
            p = next((x for x in r["products"] if x.categoryId == s.category_id), None)
            parts.append(f"{p.categoryName if p else s.category_id}({s.category_id}) n{s.n} 합{s.score_sum:+d} → 편향 {s.bias:+.2f}")
        print("   점수: " + " | ".join(parts))
    else:
        print("   점수: 리뷰 없음(또는 카탈로그에 없는 상품뿐) → 편향 0, 조회수순")
    dis = ", ".join(n for _, n in u.disliked) or "없음"
    tg = " · ".join(f"{t.axis}:{t.text}" for t in r["tags"]) or "[]"
    print(f"   태그: 선호 [{tg}] · 비선호 [{dis}]  ({r['prompt_version'] or '태그 추출 꺼짐 또는 입력 없음'})")
    for w in r["warnings"]:
        print(f"   서버 경고: {w[:140]}")
    print(f"   고르기: {r['rule']}")
    for k, (rk, s) in enumerate(r["pool"][:top], start=1):
        p = s.product
        pen = f" {s.penalty:+.2f}" if s.penalty else ""
        print(f"   {k:>2}. #{p.productId} [{p.parentCategoryName}>{p.categoryName}] {p.name[:28]}  총 {s.total:.3f} = 인기 {s.popularity:.3f} + 편향 {s.bias:+.3f}{pen}"
              f"  (점수 순위 {rk})")
    for c in checks:
        if not c[2]:
            print(f"   ❌ {c[0]} · {c[1]} — {c[3]}")


def write_files(results, checks_by, out_dir: Path, scenarios: Path) -> tuple[Path, Path]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().astimezone().strftime("%Y%m%d_%H%M%S")   # 파일 이름은 실행한 사람의 현지 시각
    xlsx, js = out_dir / f"시나리오_결과_{stamp}.xlsx", out_dir / f"시나리오_결과_{stamp}.json"
    font, hfont, hfill = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="007D77")
    red, green = PatternFill("solid", fgColor="F4CCCC"), PatternFill("solid", fgColor="D9EAD3")
    wb = Workbook()

    def sheet(title, head, rows, widths, first=False):
        ws = wb.active if first else wb.create_sheet(title)
        ws.title = title
        ws.append(head)
        for c in ws[1]:
            c.font, c.fill = hfont, hfill
        for row in rows:
            ws.append(row)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font, c.alignment = font, Alignment(vertical="top", wrap_text=isinstance(c.value, str) and len(c.value) > 40)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        return ws

    summ = []
    for r in results:
        ch = checks_by[(r["user"].no, r["mode"])]
        auto, exp, obs = [c for c in ch if c[0] == "자동"], [c for c in ch if c[0] == "기대·필수"], [c for c in ch if c[0] == OBSERVE]
        ok = all(c[2] for c in ch if blocking(c))
        summ.append([r["user"].no, r["user"].alias, r["mode"], r["rid"], len(r["user"].reviews), len(r["scores"]),
                     len(r["calls"][0]["body"]["recommendedProductIds"]) if r["calls"] else 0,
                     round(r["t202"] * 1000, 2), round(r["t_cb"] * 1000, 2) if r["t_cb"] else None,
                     round(float((r["run"] or {}).get("srv") or 0) * 1000, 2), f"{sum(c[2] for c in auto)}/{len(auto)}",
                     f"{sum(c[2] for c in exp)}/{len(exp)}", f"{sum(c[2] for c in obs)}/{len(obs)}", "PASS" if ok else "FAIL", r["user"].note])
    ws = sheet("요약", ["사용자", "별칭", "모드", "수신자 번호", "리뷰 수", "점수 소분류 수", "콜백 상품 수", "202(ms)", "7.7 콜백까지(ms)",
                       "서버 처리(ms)", "자동 검사", "기대(필수)", "기대(관찰)", "결과", "보고 싶은 것"], summ, [7, 16, 9, 12, 7, 9, 9, 9, 11, 10, 9, 10, 10, 7, 60], first=True)
    for row in ws.iter_rows(min_row=2):
        row[12].fill = green if row[12].value == "PASS" else red

    srows = []
    for r in results:
        if r["mode"] != results[0]["mode"]:
            continue
        name_of = {p.categoryId: (p.categoryName, p.parentCategoryName) for p in r["products"]}
        for s in r["scores"]:
            rv = ", ".join(f"#{x['productId']}·{x['rating']}점" for x in r["user"].reviews if x["productId"] in {p.productId for p in r["products"] if p.categoryId == s.category_id})
            srows.append([r["user"].no, r["user"].alias, s.category_id, name_of.get(s.category_id, ("", ""))[0], name_of.get(s.category_id, ("", ""))[1],
                          s.n, s.score_sum, round(s.bias, 4), rv])
        if not r["scores"]:
            srows.append([r["user"].no, r["user"].alias, None, "(점수 없음 — 편향 0)", None, 0, 0, 0.0, ""])
    sheet("점수_책정", ["사용자", "별칭", "소분류 ID", "소분류", "대분류", "n(리뷰 수)", "점수 합(별점−3)", "편향 = 합 ÷ (n+2)", "근거 리뷰"], srows,
          [7, 16, 9, 16, 12, 9, 12, 14, 40])

    rec, dis_rows = [], []
    for r in results:
        reviewed = {s.category_id for s in r["scores"]}
        for k, (rk, s) in enumerate(r["pool"], start=1):
            p = s.product
            rec.append([r["user"].no, r["mode"], k, p.productId, p.name, p.brand, p.parentCategoryName, p.categoryName, p.viewCount,
                        round(s.popularity, 4), round(s.bias, 4), round(s.penalty, 4), round(s.total, 4), "예" if s.disliked else "",
                        "예" if p.categoryId in reviewed else "", rk])
        if r["mode"] == "penalty" and r["user"].disliked:
            best: dict[int, tuple[int, Any]] = {}
            for k, s in enumerate(r["ranked"], start=1):
                if s.disliked and s.product.categoryId not in best:
                    best[s.product.categoryId] = (k, s)
            for cid, (k, s) in sorted(best.items(), key=lambda kv: kv[1][0])[:15]:
                dis_rows.append([r["user"].no, r["user"].alias, s.product.parentCategoryName, s.product.categoryName, k, len(r["ranked"]),
                                 round(s.total, 4), round(s.bias, 4), "30위 안" if k <= 30 else "30위 밖"])
    sheet("추천30", ["사용자", "모드", "순위", "상품 ID", "상품명", "브랜드", "대분류", "소분류", "조회수", "인기(0~2)", "편향", "비선호 감점", "총점",
                    "비선호", "리뷰한 소분류", "점수 순위(전체)"], rec, [7, 9, 6, 9, 44, 16, 12, 16, 9, 9, 8, 9, 8, 7, 9, 10])
    sheet("비선호_위치", ["사용자", "별칭", "대분류", "소분류", "전체 순위(최고)", "후보 수", "총점", "편향", "판정"], dis_rows, [7, 16, 12, 16, 12, 9, 8, 8, 9])

    tags = [[r["user"].no, r["mode"], " · ".join(f"{t.axis}:{t.text}" for t in r["tags"]) or "[]",
             json.dumps([n for _, n in r["user"].disliked], ensure_ascii=False), r["prompt_version"] or "태그 추출 꺼짐 또는 입력 없음",
             r["user"].preference or ""] for r in results]
    sheet("태그", ["사용자", "모드", "preferred_tags(축:태그)", "disliked_tags", "프롬프트@모델", "취향 문장(입력, 저장 안 함)"], tags, [7, 9, 50, 30, 30, 40])
    sheet("콜백_원문", ["사용자", "모드", "경로", "Authorization", "본문"], [[r["user"].no, r["mode"], c["path"], c["auth"], json.dumps(c["body"], ensure_ascii=False)]
                                                                         for r in results for c in r["calls"]], [7, 9, 44, 24, 120])
    crow = [[no, mode, k, n, "통과" if ok else "실패", d] for (no, mode), ch in checks_by.items() for k, n, ok, d in ch]
    ws = sheet("검사", ["사용자", "모드", "종류", "검사", "결과", "내용"], crow, [7, 9, 10, 60, 7, 70])
    for row in ws.iter_rows(min_row=2):
        row[4].fill = green if row[4].value == "통과" else red
    inp = [[u.no, u.alias, u.persona, "; ".join(f"{c}:{n}" for c, n in u.disliked), u.preference or "",
            " | ".join(f"#{x['productId']} {x['rating']}점" for x in u.reviews)] for u in {r["user"].no: r["user"] for r in results}.values()]
    sheet("입력", ["사용자", "별칭", "페르소나", "비선호", "취향 문장", f"리뷰 (원본: {scenarios.name})"], inp, [7, 16, 40, 30, 36, 60])
    trows = [[r["user"].no, r["user"].alias, r["mode"], r["config"]["model"], len(r["user"].reviews), "y" if r["user"].preference else "n",
              *(r["usage"] or [None] * 8)] for r in results if r["config"]["tagger"] != "none"]
    if trows:
        sheet("토큰", ["사용자", "별칭", "모드", "모델", "리뷰 수", "취향 문장", "입력 토큰", "출력 토큰", "LLM 전체(ms)", "적재(ms)", "입력 처리(ms)",
                     "출력 생성(ms)", "멈춘 이유", "비용(USD)"], trows, [7, 16, 9, 14, 7, 8, 9, 9, 11, 9, 11, 11, 9, 10])
    wb.save(xlsx)

    js.write_text(json.dumps([{
        "user": r["user"].no, "alias": r["user"].alias, "mode": r["mode"], "recipientUserId": r["rid"],
        "timing_ms": {"accept_202": round(r["t202"] * 1000, 3), "callback": round(r["t_cb"] * 1000, 3) if r["t_cb"] else None,
                      "server": round(float((r["run"] or {}).get("srv") or 0) * 1000, 3)},
        "category_scores": [{"category_id": s.category_id, "n": s.n, "score_sum": s.score_sum, "bias": s.bias} for s in r["scores"]],
        "callback": r["calls"][0]["body"] if r["calls"] else None,
        "rule": r["rule"], "prompt_version": r["prompt_version"], "tags": [[t.axis, t.text] for t in r["tags"]],
        "top30": [{"rank": k, "score_rank": rk, "productId": s.product.productId, "total": s.total, "popularity": s.popularity, "bias": s.bias,
                   "penalty": s.penalty} for k, (rk, s) in enumerate(r["pool"], start=1)],
        "checks": [{"kind": k, "name": n, "ok": ok, "detail": d} for k, n, ok, d in checks_by[(r["user"].no, r["mode"])]],
        "usage": r["usage"]._asdict() if r.get("usage") else None,
        "warnings": r["warnings"],
    } for r in results], ensure_ascii=False, indent=1), encoding="utf-8")
    return xlsx, js


# ---------------------------------------------------------------------------


def cleanup(engine: sa.Engine) -> None:
    with engine.begin() as c:
        for base in RID_BASE.values():
            c.execute(sa.text("delete from ai_profile.recipient_profiles where recipient_user_id between :a and :b"), {"a": base, "b": base + 9_999})
            c.execute(sa.text("delete from ai_profile.profile_runs where recipient_user_id between :a and :b"), {"a": base, "b": base + 9_999})


GOLDEN = HERE / "golden_d.json"


def _golden(results, checks_by, save: bool) -> None:
    """태그 끔 · D 결과를 기준 파일과 비교(관찰) — 다르면 의도된 변경인지 사람이 본다. save 면 지금 결과로 기준을 새로 쓴다."""
    cur = {f"{r['user'].no}|{r['mode']}": (r["calls"][0]["body"]["recommendedProductIds"] if r["calls"] else []) for r in results
           if r["config"]["tagger"] == "none" and r["config"]["strategy"] == "d"}
    if not cur:
        return
    if save:
        GOLDEN.write_text(json.dumps({"note": "태그 끔 · D · 감점/제외 모드의 콜백 30개 — 변경 시 차이를 사람이 확인(10-12 처음 저장)",
                                      "pools": cur}, ensure_ascii=False, indent=1))
        return
    if not GOLDEN.exists():
        return
    gold = json.loads(GOLDEN.read_text())["pools"]
    for r in results:
        k = f"{r['user'].no}|{r['mode']}"
        if k in cur and k in gold:
            same = len(set(cur[k]) & set(gold[k]))
            checks_by[(r["user"].no, r["mode"])].append((OBSERVE, "[회귀] 기준 파일(golden_d.json)과 같은 30개 — 다르면 의도된 변경인지 확인",
                                                         cur[k] == gold[k], f"겹침 {same}/30 · 순서 {'같음' if cur[k] == gold[k] else '다름'}"))


def run(scenarios: Path = DEFAULT_SCENARIOS, out_dir: Path = DEFAULT_OUT, modes: tuple[str, ...] = MODES, *, keep: bool = False,
        top: int = 5, quiet: bool = False, model_env: dict[str, str] | None = None, save_golden: bool = False) -> dict[str, Any]:
    """시나리오 전체 실행 → {'ok', 'xlsx', 'json', 'results', 'checks'}. pytest(tests/integration/test_scenario_callback30.py)도 이걸 부른다."""
    from profiling.settings import get_settings

    users = load_scenarios(scenarios)
    engine = sa.create_engine(get_settings().DATABASE_URL, future=True)
    cleanup(engine)
    t_all = time.perf_counter()
    results: list[dict[str, Any]] = []
    try:
        for mode in modes:
            results += run_mode(mode, users, engine, model_env)
        by_key = {(r["user"].no, r["mode"]): r for r in results}
        checks_by = {}
        for r in results:
            other = by_key.get((r["user"].no, "exclude" if r["mode"] == "penalty" else "penalty"))
            checks_by[(r["user"].no, r["mode"])] = check_run(r, engine, other)
        _golden(results, checks_by, save_golden)
        if not quiet:
            print(f"\n=== 프로파일링 코어 30개 상품 콜백 검증 — 사용자 {len(users)}명 × 모드 {len(modes)} ({', '.join(modes)}) · {scenarios.name} ===")
            for r in results:
                print_run(r, checks_by[(r["user"].no, r["mode"])], top)
        xlsx, js = write_files(results, checks_by, out_dir, scenarios)
    finally:
        if not keep:
            cleanup(engine)
        engine.dispose()
    allc = [c for ch in checks_by.values() for c in ch]
    ok = all(c[2] for c in allc if blocking(c))
    obs = [c for c in allc if not blocking(c)]
    if not quiet:
        cb = sorted(r["t_cb"] for r in results if r["t_cb"])
        blk = [c for c in allc if blocking(c)]
        print(f"\n{'✅ 전부 통과' if ok else '❌ 실패 있음'} — 검사(자동 · 필수) {sum(c[2] for c in blk)}/{len(blk)}"
              + (f" · 관찰 {sum(c[2] for c in obs)}/{len(obs)}" if obs else "") + f" · 실행 {len(results)}건 · "
              f"7.7 콜백까지 p90 {ms(cb[int(0.9 * (len(cb) - 1))]) if cb else '-'} · p95 {ms(cb[int(0.95 * (len(cb) - 1))]) if cb else '-'} · "
              f"최대 {ms(cb[-1]) if cb else '-'} · 전체 {time.perf_counter() - t_all:.1f}s")
        if (us := usage_summary(results)):
            print(f"LLM 사용량 — 호출 {us['calls']}회 · 입력 {us['input_tokens']:,} · 출력 {us['output_tokens']:,} 토큰 "
                  f"(사람당 입력 p50 {us['input_p50']} · p95 {us['input_p95']} · 출력 p50 {us['output_p50']} · p95 {us['output_p95']}) · "
                  f"LLM p95 {us['llm_ms_p95']}ms · 멈춤 {us['stops']}" + (f" · 비용 ${us['cost_usd']}" if us["cost_usd"] is not None else ""))
        print(f"결과 파일: {xlsx}\n          {js}")
    return {"ok": ok, "xlsx": xlsx, "json": js, "results": results, "checks": checks_by}


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--scenarios", type=Path, default=DEFAULT_SCENARIOS)
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    ap.add_argument("--modes", default=",".join(MODES), help="exclude,penalty 중 쉼표로")
    ap.add_argument("--keep", action="store_true", help="실행 뒤 시험 수신자 행을 지우지 않음(DB 로 직접 보기)")
    ap.add_argument("--top", type=int, default=5, help="화면에 보일 상위 상품 수")
    ap.add_argument("--tagger", choices=("none", "ollama", "openrouter"), default="none",
                    help="태그 추출(LLM) — ollama: 로컬 Ollama · openrouter: 외부 API(유료, profiling/.env 의 PROFILING_OPENROUTER_API_KEY)")
    ap.add_argument("--tagger-model", default="gemma4:12b")
    ap.add_argument("--embedding-dir", default=None, help="product-search 폴더 — 주면 D 의 임베딩 단계가 켜진다")
    ap.add_argument("--save-golden", action="store_true", help="태그 끔 · D 결과로 기준 파일(golden_d.json)을 새로 쓴다 — 의도된 변경 뒤에만")
    a = ap.parse_args()
    model_env = {"PROFILING_TAGGER": a.tagger, "PROFILING_TAGGER_MODEL": a.tagger_model}
    if a.embedding_dir:
        model_env["PROFILING_EMBEDDING_DIR"] = a.embedding_dir
    modes = tuple(m.strip() for m in a.modes.split(",") if m.strip())
    return 0 if run(a.scenarios, a.out, modes, keep=a.keep, top=a.top, model_env=model_env, save_golden=a.save_golden)["ok"] else 1


if __name__ == "__main__":
    sys.exit(main())
