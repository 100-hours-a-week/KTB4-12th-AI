"""태그 추출 모델 비교 — 같은 시나리오를 추출 모델마다(claude-opus-5-5 · gemma4:12b · qwen3.5:9b …) 임베딩 비교에 돌려 나란히 본다.

    cd profiling && uv run python -m tools.embed_compare.compare_taggers [--taggers claude-opus-5-5,gemma4:12b,qwen3.5:9b]

비교하는 것(기준 = 첫 번째 추출 모델)
  태그     개수 · 축 분포 · 기준과 **정확 일치**(공백 무시) · **근접 일치**(같은 임베딩 모델로 태그 글자를 벡터로 바꿔 코사인 ≥ SOFT)
  추천     D 30개의 소분류 종수 · 최다 · 편향 소분류 · 비선호 · 조회수 · **기준 D 30개와 겹침**
  시간     추출 모델의 사람당 추출 시간(out/태그추출_<모델>_*.json 의 마지막 기록) — 기준(claude-opus-5-5 수작업)은 기록 없음
결과: 화면 + out/태그추출_비교_<시각>.xlsx · .json. 감점 모드만(D 는 두 모드 결과가 같다).
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime
from typing import Any

import numpy as np

from tools.embed_compare import run as ec

SOFT = 0.80


def _norm(t: str) -> str:
    return re.sub(r"\s+", "", t)


def extraction_log(tagger: str) -> dict[str, Any] | None:
    logs = sorted(ec.DEFAULT_OUT.glob(f"태그추출_{re.sub(r'[^\w.-]', '_', tagger)}_*.json"))
    return json.loads(logs[-1].read_text()) if logs else None


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--taggers", default="claude-opus-5-5,gemma4:12b,qwen3.5:9b")
    a = ap.parse_args()
    taggers = [t.strip() for t in a.taggers.split(",") if t.strip()]
    base = taggers[0]
    runs = {t: ec.run(modes=("penalty",), quiet=True, tagger=t) for t in taggers}
    tags = {t: ec.load_tags(ec.DEFAULT_SCENARIOS, t) for t in taggers}
    users = [r["user"].no for r in runs[base]["results"]]
    alltags = sorted({tg for t in taggers for v in tags[t].values() for _, tg in v})
    with ec.QueryEncoder() as enc:
        vec = {tg: enc(tg) for tg in alltags}

    rows_tag, rows_rec, summary = [], [], []
    for t in taggers:
        log = extraction_log(t)
        exact = soft = nbase = 0
        for u in users:
            mine, ref = tags[t].get(u, []), tags[base].get(u, [])
            nbase += len(ref)
            for _, rt in ref:
                exact += any(_norm(rt) == _norm(mt) for _, mt in mine)
                soft += any(float(vec[rt] @ vec[mt]) >= SOFT for _, mt in mine)
            sec = log["users"][u]["seconds"] if log and u in log["users"] else None
            rows_tag.append((t, u, len(mine), " · ".join(f"{ax}:{tg}" for ax, tg in mine) or "-", sec))
        res = {r["user"].no: r for r in runs[t]["results"]}
        bres = {r["user"].no: r for r in runs[base]["results"]}
        for u in users:
            m, r = res[u]["m"]["D"], res[u]
            ov = len({r["ranked"][i].product.productId for i in r["st"]["D"]} & {bres[u]["ranked"][i].product.productId for i in bres[u]["st"]["D"]})
            rows_rec.append((t, u, len(tags[t].get(u, [])), m["소분류_종수"], m["최다_개수"], m["편향_소분류"], m["비선호"], m["조회수_백분위_평균"], ov))
        axes: dict[str, int] = {}
        for v in tags[t].values():
            for ax, _ in v:
                axes[ax] = axes.get(ax, 0) + 1
        secs = sorted(x for x in (log["users"][u]["seconds"] for u in users) if x is not None) if log else []
        tagged = [x for x in rows_rec if x[0] == t and x[2]]
        summary.append({"tagger": t, "digest": (log or {}).get("digest", "")[:12], "tags": sum(len(v) for v in tags[t].values()),
                        "users_with_tags": sum(1 for u in users if tags[t].get(u)), "axes": axes,
                        "exact": f"{exact}/{nbase}", "soft": f"{soft}/{nbase}",
                        "p95_s": round(float(np.percentile(secs, 95)), 2) if secs else None, "max_s": secs[-1] if secs else None,
                        "d_species_avg": round(float(np.mean([x[3] for x in tagged])), 1) if tagged else None,
                        "d_overlap_avg": round(float(np.mean([x[8] for x in tagged])), 1) if tagged else None,
                        "d_checks_ok": runs[t]["ok"]})

    print(f"\n=== 태그 추출 모델 비교 — 기준 {base} · 임베딩 {runs[base]['vectors']['model']} · 근접 일치 코사인 ≥ {SOFT} ===")
    for s in summary:
        print(f"  {s['tagger']:<16} 태그 {s['tags']:>2}(사람 {s['users_with_tags']}) · 축 {s['axes']} · 기준 일치 정확 {s['exact']} · 근접 {s['soft']} · "
              f"추출 p95 {s['p95_s'] or '-'}s · D 종수 평균 {s['d_species_avg']} · 기준 D 겹침 평균 {s['d_overlap_avg']}/30 · D 검사 {'통과' if s['d_checks_ok'] else '실패'}")
    for t in taggers:
        print(f"\n  [{t}]")
        for _, u, n, txt, sec in (x for x in rows_tag if x[0] == t):
            rr = next(x for x in rows_rec if x[0] == t and x[1] == u)
            print(f"   {u} 태그 {n} {txt[:70]:<70} | D {rr[3]:>2}종/{rr[4]:>2} 비선호 {rr[6]} 겹침 {rr[8]:>2}" + (f" | {sec:.1f}s" if sec else ""))

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    stamp = datetime.now().astimezone().strftime("%Y%m%d_%H%M%S")
    xlsx, js = ec.DEFAULT_OUT / f"태그추출_비교_{stamp}.xlsx", ec.DEFAULT_OUT / f"태그추출_비교_{stamp}.json"
    wb = Workbook()
    hfont, hfill = Font(name="Arial", size=10, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="007D77")

    def sheet(title, head, rows, first=False):
        ws = wb.active if first else wb.create_sheet(title)
        ws.title = title
        ws.append(head)
        for c in ws[1]:
            c.font, c.fill = hfont, hfill
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "A2"

    sheet("요약", ("추출 모델", "Ollama digest", "태그 수", "태그 있는 사람", "축 분포", "기준 일치(정확)", f"기준 일치(근접 ≥{SOFT})",
                 "추출 p95(s)", "추출 최대(s)", "D 종수 평균", "기준 D 겹침 평균(/30)", "D 검사"),
          [(s["tagger"], s["digest"], s["tags"], s["users_with_tags"], json.dumps(s["axes"], ensure_ascii=False), s["exact"], s["soft"],
            s["p95_s"], s["max_s"], s["d_species_avg"], s["d_overlap_avg"], "통과" if s["d_checks_ok"] else "실패") for s in summary], first=True)
    sheet("태그", ("추출 모델", "번호", "태그 수", "태그", "추출 시간(s)"), rows_tag)
    sheet("D_추천", ("추출 모델", "번호", "태그 수", "소분류 종수", "최다", "편향 소분류", "비선호", "조회수 백분위", "기준 D 겹침(/30)"), rows_rec)
    for w, col in zip((16, 14, 8, 10, 30, 12, 14, 10, 10, 10, 14, 8), "ABCDEFGHIJKL"):
        wb["요약"].column_dimensions[col].width = w
    wb["태그"].column_dimensions["D"].width = 80
    wb.save(xlsx)
    js.write_text(json.dumps({"base": base, "soft": SOFT, "summary": summary, "tags": rows_tag, "recommend": rows_rec}, ensure_ascii=False, indent=1))
    print(f"\n결과 파일: {xlsx}\n          {js}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
