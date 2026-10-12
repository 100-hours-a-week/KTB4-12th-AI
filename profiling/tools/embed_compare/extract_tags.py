"""로컬 LLM(Ollama)으로 시나리오 사용자의 선호 태그를 뽑아 '태그' 시트에 그 모델 이름으로 적는다 — Claude 수작업 태그와 비교용.

    cd profiling && uv run python -m tools.embed_compare.extract_tags --models gemma4:12b,qwen3.5:9b
    uv run python -m tools.embed_compare.run --tagger gemma4:12b
    uv run python -m tools.embed_compare.extract_tags --from-json out/claude_prompt_responses.json   # 같은 프롬프트에 대한 다른 모델의 원출력 넣기

입력은 실제 서비스에서 AI 가 받는 것만 쓴다: 취향 문장(giftPreference) · 리뷰(별점 · 글) + 리뷰 상품의 카탈로그 정보(이름 · 브랜드 ·
대분류 > 소분류 · product_kind). 페르소나 · '보고 싶은 것'(시나리오 설명)은 넣지 않는다. 비선호는 코드가 따로 처리하므로 넣지 않는다.
같은 상품 중복 리뷰는 앞(최신)만, 카탈로그에 없는 상품은 빼고 넘긴다(코어 규칙과 같음).

규칙은 팀 폴더 '태그추출_프롬프트_유의사항.md' 1 · 2 · 3 · 6절(판촉 단어 · 증정품 · 구매 경험 · 4축)을 옮겼다. 결정성: temperature 0 · seed 고정 ·
생각 모드 끔 · JSON 스키마 강제. 검증기: 축이 4개 중 하나 · 근거 ref(여러 개 가능)가 실제 리뷰/취향 · 근거가 별점 1~3 리뷰뿐이면 버림 · 판촉 단어 태그는 버림.
기록: out/태그추출_<모델>_<시각>.json(프롬프트 해시 · 모델 digest · 사람마다 시간 · 원출력 · 버린 태그와 사유).
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Any

import httpx
import sqlalchemy as sa

from profiling import tagging
from profiling.schemas import ProductRecord
from profiling.types import ProfileRequest, Review
from tools.embed_compare.run import DEFAULT_OUT, DEFAULT_SCENARIOS

OLLAMA = "http://localhost:11434"
SEED = 42
# 프롬프트 · 출력 형식 · 검증기는 코어 profiling.tagging 이 정본이다(10-12 코어 반영) — 도구와 서비스가 한 프롬프트(해시 2fc950d4)를 쓴다.
AXES, PROMO, SYSTEM, SCHEMA, validate = tagging.AXES, tagging.PROMO, tagging.SYSTEM, tagging.SCHEMA, tagging.validate


def user_prompt(u, info: dict[int, ProductRecord]) -> tuple[str, dict[str, int]]:
    """시나리오 사용자 하나 → (프롬프트, ref → 별점). 코어 tagging.build_prompt 그대로."""
    rq = ProfileRequest(recipient_user_id=1, source_version=1, gift_preference=u.preference, disliked_categories=[],
                        reviews=[Review(product_id=r["productId"], rating=r["rating"], review_text=r["reviewText"]) for r in u.reviews])
    return tagging.build_prompt(rq, info)


def extract(model: str, users, info) -> dict[str, Any]:
    digest = next((m["digest"] for m in httpx.get(f"{OLLAMA}/api/tags", timeout=10).json()["models"] if m["name"] == model), None)
    if digest is None:
        raise SystemExit(f"Ollama 에 {model} 이 없음 — ollama pull {model}")
    out = {"model": model, "digest": digest, "seed": SEED, "prompt_sha256": hashlib.sha256((SYSTEM + json.dumps(SCHEMA)).encode()).hexdigest(),
           "users": {}}
    with httpx.Client(base_url=OLLAMA, timeout=600) as c:
        c.post("/api/generate", json={"model": model, "prompt": "", "keep_alive": "10m"})          # 모델 올리기(시간 측정에서 뺌)
        for u in users:
            prompt, ratings = user_prompt(u, info)
            t0 = time.perf_counter()
            r = c.post("/api/chat", json={"model": model, "stream": False, "think": False, "format": SCHEMA,
                                          "options": {"temperature": 0, "seed": SEED},
                                          "messages": [{"role": "system", "content": SYSTEM}, {"role": "user", "content": prompt}]})
            dt = time.perf_counter() - t0
            r.raise_for_status()
            body = r.json()
            raw = body["message"]["content"]
            try:
                tags = json.loads(raw).get("tags", [])
            except json.JSONDecodeError:
                tags = []
            keep, drop = validate(tags, ratings, bool(u.preference))
            out["users"][u.no] = {"seconds": round(dt, 3), "eval_tokens": body.get("eval_count"), "prompt": prompt, "raw": raw,
                                  "tags": keep, "dropped": drop}
            print(f"  {model} {u.no} {dt:5.1f}s · 태그 {len(keep)} · 버림 {len(drop)} | " + " · ".join(f"{t['axis']}:{t['tag']}" for t in keep))
    return out


def extract_openrouter(model: str, users, info, repeat: int = 1) -> dict[str, Any]:
    """외부 API(OpenRouter) — 코어 어댑터 OpenRouterTagExtractor 로 같은 SYSTEM · 스키마 · temperature 0 · seed 42. **유료.**
    repeat 번 반복해 시간 분포와 안정성(반복마다 같은 태그인가)을 본다. 태그 · 원출력은 첫 회 기준. 키는 profiling/.env."""
    from profiling.llm import OpenRouterTagExtractor
    from profiling.ports import TagExtractError
    from profiling.settings import get_settings

    st = get_settings()
    if st.OPENROUTER_API_KEY is None or not st.OPENROUTER_API_KEY.get_secret_value().strip():
        raise SystemExit("PROFILING_OPENROUTER_API_KEY 없음 — profiling/.env 에 넣기")
    tagger = OpenRouterTagExtractor(st.OPENROUTER_URL, st.OPENROUTER_API_KEY.get_secret_value(), model, timeout_s=st.TAGGER_TIMEOUT_S, seed=SEED)
    out = {"model": model, "digest": "openrouter", "seed": SEED, "repeat": repeat,
           "prompt_sha256": hashlib.sha256((SYSTEM + json.dumps(SCHEMA)).encode()).hexdigest(), "users": {}}
    try:
        for u in users:
            prompt, ratings = user_prompt(u, info)
            runs = []
            for _ in range(repeat):
                try:
                    reply = tagger.extract(SYSTEM, prompt, SCHEMA)
                except TagExtractError as e:
                    runs.append({"error": e.kind, "usage": e.usage._asdict() if e.usage else None})
                    continue
                keep, drop = validate(tagging.parse(reply.text), ratings, bool(u.preference))
                runs.append({"raw": reply.text, "keep": keep, "drop": drop, "usage": reply.usage._asdict() if reply.usage else None})
            ok = [r for r in runs if "error" not in r]
            first = ok[0] if ok else {"raw": "", "keep": [], "drop": []}
            sets = [sorted((t["axis"], t["tag"]) for t in r["keep"]) for r in ok]
            us = [r["usage"] for r in ok if r.get("usage")]
            out["users"][u.no] = {
                "seconds": round(us[0]["total_ms"] / 1000, 3) if us else None,
                "seconds_all": [round(x["total_ms"] / 1000, 3) for x in us],
                "eval_tokens": us[0]["output_tokens"] if us else None,
                "input_tokens": [x["input_tokens"] for x in us], "output_tokens": [x["output_tokens"] for x in us],
                "cost_usd": [x["cost_usd"] for x in us], "errors": [r["error"] for r in runs if "error" in r],
                "stable": len({json.dumps(x, ensure_ascii=False) for x in sets}) <= 1,
                "prompt": prompt, "raw": first["raw"], "tags": first["keep"], "dropped": first["drop"]}
            d = out["users"][u.no]
            print(f"  {model} {u.no} {d['seconds_all']}s · 태그 {len(d['tags'])} · 버림 {len(d['dropped'])} · 안정 {d['stable']}"
                  f"{' · 오류 ' + str(d['errors']) if d['errors'] else ''} | " + " · ".join(f"{t['axis']}:{t['tag']}" for t in d["tags"]))
    finally:
        tagger.close()
    return out


def from_responses(path: Path, users, info) -> dict[str, Any]:
    """Ollama 밖에서 같은 SYSTEM · user_prompt 로 받은 원출력(JSON)을 같은 검증기로 — 예: claude-opus-5-5 가 대화 안에서 답한 것.
    파일 모양: {"model": 이름, "note": 설명, "responses": {"U01": {"tags": [...]}, ...}}. 시간은 없다(API 호출이 아님)."""
    doc = json.loads(path.read_text())
    out = {"model": doc["model"], "digest": "", "seed": None, "note": doc.get("note", ""),
           "prompt_sha256": hashlib.sha256((SYSTEM + json.dumps(SCHEMA)).encode()).hexdigest(), "users": {}}
    for u in users:
        prompt, ratings = user_prompt(u, info)
        raw = doc["responses"].get(u.no, {"tags": []})
        keep, drop = validate(raw.get("tags", []), ratings, bool(u.preference))
        out["users"][u.no] = {"seconds": None, "eval_tokens": None, "prompt": prompt, "raw": json.dumps(raw, ensure_ascii=False),
                              "tags": keep, "dropped": drop}
        print(f"  {doc['model']} {u.no} 태그 {len(keep)} · 버림 {len(drop)} | " + " · ".join(f"{t['axis']}:{t['tag']}" for t in keep)
              + ("".join(f" ✗{t['tag']}({t['drop']})" for t in drop)))
    return out


def write_sheet(path: Path, model: str, res: dict[str, Any]) -> None:
    """'태그' 시트에서 그 모델 행을 지우고 새로 적는다. 다른 모델(claude-opus-5-5 등) 행은 건드리지 않는다."""
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb["태그"]
    for r in range(ws.max_row, 1, -1):
        if ws.cell(r, 5).value == model:
            ws.delete_rows(r)
    how = (f"tools/embed_compare/extract_tags.py · OpenRouter(외부 API) · temperature 0 · seed {SEED} · 반복 {res.get('repeat')} · 프롬프트 {res['prompt_sha256'][:8]} · {datetime.now().astimezone():%m-%d}"
           if res.get("digest") == "openrouter" else
           f"tools/embed_compare/extract_tags.py · Ollama · temperature 0 · seed {SEED} · 생각 끔 · 프롬프트 {res['prompt_sha256'][:8]} · {datetime.now().astimezone():%m-%d}"
           if res.get("seed") is not None else f"같은 프롬프트 {res['prompt_sha256'][:8]} · {res.get('note', '')[:80]} · {datetime.now().astimezone():%m-%d}")
    for no, u in res["users"].items():
        for t in u["tags"]:
            ws.append((no, t["axis"], t["tag"], f"{t['ref']}: {t['evidence']}"[:120], model, how))
    wb.save(path)


def main() -> int:
    from profiling.settings import get_settings
    from tools.scenario_callback30.run import load_scenarios

    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--models", default="gemma4:12b,qwen3.5:9b")
    ap.add_argument("--scenarios", type=Path, default=DEFAULT_SCENARIOS)
    ap.add_argument("--openrouter", action="store_true", help="Ollama 대신 OpenRouter(유료) — --models 에 OpenRouter 모델 id(예 anthropic/claude-haiku-5.5)")
    ap.add_argument("--repeat", type=int, default=1, help="--openrouter 일 때 사람마다 반복 횟수(시간 분포 · 안정성)")
    ap.add_argument("--from-json", type=Path, help="Ollama 대신 원출력 파일(같은 프롬프트에 대한 다른 모델의 답)을 검증해 시트에 넣음")
    a = ap.parse_args()
    users = load_scenarios(a.scenarios)
    eng = sa.create_engine(get_settings().DATABASE_URL, future=True)
    from profiling.catalog import DbCatalogReader
    _, products = DbCatalogReader(eng).active()
    info = {p.productId: p for p in products}
    eng.dispose()
    DEFAULT_OUT.mkdir(parents=True, exist_ok=True)
    if a.from_json:
        res = from_responses(a.from_json, users, info)
        write_sheet(a.scenarios, res["model"], res)
        log = DEFAULT_OUT / f"태그추출_{re.sub(r'[^\w.-]', '_', res['model'])}_{datetime.now().astimezone():%Y%m%d_%H%M%S}.json"
        log.write_text(json.dumps(res, ensure_ascii=False, indent=1))
        print(f"  → 태그 {sum(len(u['tags']) for u in res['users'].values())} · 버림 {sum(len(u['dropped']) for u in res['users'].values())} · 기록 {log.name}")
        return 0
    for model in (m.strip() for m in a.models.split(",") if m.strip()):
        print(f"\n=== 태그 추출 — {model}")
        res = extract_openrouter(model, users, info, a.repeat) if a.openrouter else extract(model, users, info)
        write_sheet(a.scenarios, model, res)
        log = DEFAULT_OUT / f"태그추출_{re.sub(r'[^\w.-]', '_', model)}_{datetime.now().astimezone():%Y%m%d_%H%M%S}.json"
        log.write_text(json.dumps(res, ensure_ascii=False, indent=1))
        secs = sorted(u["seconds"] for u in res["users"].values())
        print(f"  → 사람당 p90 {secs[int(0.9 * (len(secs) - 1))]:.1f}s · 최대 {secs[-1]:.1f}s · 태그 {sum(len(u['tags']) for u in res['users'].values())} · "
              f"버림 {sum(len(u['dropped']) for u in res['users'].values())} · 기록 {log.name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
