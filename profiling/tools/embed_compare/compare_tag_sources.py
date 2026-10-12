"""출처별 태그 비교표 — 같은 리뷰 · 같은 취향 문장에서 모델마다 어떤 태그를 뽑았나를 한 행에 나란히.

    cd profiling && uv run python -m tools.embed_compare.compare_tag_sources [--models claude-opus-5-5-prompt,gemma4:12b,qwen3.5:9b,claude-opus-5-5]

행 = (사람, 출처) — 출처는 취향 문장(pref) 또는 리뷰 r1, r2 …(같은 상품 중복 · 카탈로그에 없는 상품은 프롬프트와 같은 규칙으로 뺌).
태그가 어느 출처에서 나왔는지는 각 모델 추출 기록(out/태그추출_<모델>_*.json 의 마지막)의 근거 ref 를 쓴다 — "pref, r1" 처럼 여럿이면 그 행들에 모두.
claude-opus-5-5(대화 중 수작업, 고정 프롬프트 없음)는 기록 대신 시나리오 엑셀 '태그' 시트의 근거 글("리뷰 399 …" · "취향 문장")에서 출처를 읽는다.
결과: out/태그_출처별_비교_<시각>.xlsx — 시트: 출처별_비교 · 요약 · 태그_전체 · 설명.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime
from pathlib import Path

from tools.embed_compare.run import DEFAULT_OUT, DEFAULT_SCENARIOS

MANUAL = "claude-opus-5-5"
_LINE = re.compile(r"^(r\d+)\. 별점 (\d)점 · 상품명: (.*?) · 브랜드: (.*?) · 분류: (.*?) · 종류: (.*?) · 리뷰 글: (.*)$")


def latest_log(model: str) -> dict | None:
    logs = sorted(DEFAULT_OUT.glob(f"태그추출_{re.sub(r'[^\w.-]', '_', model)}_2*.json"))
    return json.loads(logs[-1].read_text()) if logs else None


def sources(prompt: str) -> list[dict]:
    """프롬프트 → 출처 목록(취향 문장 + 리뷰). 모든 모델이 같은 프롬프트를 받았으므로 어느 기록에서 읽어도 같다."""
    out = []
    first = prompt.splitlines()[0]
    out.append({"ref": "pref", "kind": "취향 문장", "rating": None, "text": first.split(": ", 1)[1] if ": " in first else "(없음)",
                "product": "", "path": "", "pid": None})
    for line in prompt.splitlines():
        m = _LINE.match(line.strip())
        if m:
            ref, rating, name, brand, path, kind, text = m.groups()
            out.append({"ref": ref, "kind": "리뷰", "rating": int(rating), "text": text, "product": f"{name} · {brand}", "path": f"{path} · {kind}",
                        "pid": None})
    return out


def manual_tags(scenarios: Path, reviews_by_user: dict[str, list[int]]) -> dict[str, list[dict]]:
    """수작업 태그의 근거 글 → ref. "리뷰 399" 는 그 사람의 리뷰 순서(중복 제거 뒤)로 r 번호를 찾는다."""
    from openpyxl import load_workbook

    out: dict[str, list[dict]] = {}
    for r in load_workbook(scenarios, read_only=True, data_only=True)["태그"].iter_rows(min_row=2, values_only=True):
        if not r or not r[0] or (r[4] if len(r) > 4 else None) not in (None, MANUAL):
            continue
        no, axis, tag, why = r[0], r[1], str(r[2]).strip(), str(r[3] or "")
        refs = ["pref"] if "취향" in why else []
        order = reviews_by_user.get(no, [])
        for pid in re.findall(r"리뷰\s*([\d ·]+)", why):
            for x in re.findall(r"\d{3,6}", pid):
                if int(x) in order:
                    refs.append(f"r{order.index(int(x)) + 1}")
        out.setdefault(no, []).append({"axis": axis, "tag": tag, "ref": ", ".join(refs) or "?", "evidence": why})
    return out


def main() -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from profiling.catalog import DbCatalogReader
    from profiling.settings import get_settings
    from tools.scenario_callback30.run import load_scenarios

    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--models", default="claude-opus-5-5-prompt,gemma4:12b,qwen3.5:9b,claude-opus-5-5")
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    a = ap.parse_args()
    models = [m.strip() for m in a.models.split(",") if m.strip()]
    logs = {m: latest_log(m) for m in models if m != MANUAL}
    base_log = next(v for v in logs.values() if v)
    users = load_scenarios(DEFAULT_SCENARIOS)

    import sqlalchemy as sa
    _, products = DbCatalogReader(sa.create_engine(get_settings().DATABASE_URL)).active()
    in_cat = {p.productId for p in products}
    order: dict[str, list[int]] = {}
    for u in users:
        seen: list[int] = []
        for rv in u.reviews:
            if rv["productId"] in in_cat and rv["productId"] not in seen:
                seen.append(rv["productId"])
        order[u.no] = seen
    tags = {m: ({no: v["tags"] for no, v in logs[m]["users"].items()} if m != MANUAL else manual_tags(DEFAULT_SCENARIOS, order)) for m in models}
    dropped = {m: ({no: v.get("dropped", []) for no, v in logs[m]["users"].items()} if m != MANUAL else {}) for m in models}

    rows, long = [], []
    for u in users:
        for src in sources(base_log["users"][u.no]["prompt"]):
            if src["kind"] == "취향 문장" and src["text"] == "(없음)":
                continue
            line = {"user": u.no, "alias": u.alias, **src, "by": {}}
            for m in models:
                hit = [t for t in tags[m].get(u.no, []) if src["ref"] in re.split(r"[\s,/·]+", str(t.get("ref", "")))]
                line["by"][m] = hit
                for t in hit:
                    shared = len([x for x in re.split(r"[\s,/·]+", str(t.get("ref", ""))) if x]) > 1
                    long.append((u.no, src["ref"], m, t["axis"], t["tag"], t.get("ref", ""), "여러 출처" if shared else ""))
            rows.append(line)

    stamp = datetime.now().astimezone().strftime("%Y%m%d_%H%M%S")
    path = a.out / f"태그_출처별_비교_{stamp}.xlsx"
    wb = Workbook()
    F, HF = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True, color="FFFFFF")
    HFILL, LOW, SAME = PatternFill("solid", fgColor="007D77"), PatternFill("solid", fgColor="F3F3F3"), PatternFill("solid", fgColor="D9EAD3")
    wrap = Alignment(wrap_text=True, vertical="top")

    def sheet(title, head, data, widths, first=False):
        ws = wb.active if first else wb.create_sheet(title)
        ws.title = title
        ws.append(head)
        for c in ws[1]:
            c.font, c.fill, c.alignment = HF, HFILL, wrap
        for r in data:
            ws.append(r)
            for c in ws[ws.max_row]:
                c.font, c.alignment = F, wrap
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "E2"
        return ws

    fmt = lambda ts: "\n".join(f"{t['axis']}:{t['tag']}" for t in ts) or "—"
    head = ["사람", "출처", "별점", "입력(취향 문장 · 리뷰 글)", "상품 · 브랜드", "분류 · 종류"] + models + ["모든 모델 공통", "한 모델에만"]
    data = []
    for r in rows:
        sets = {m: {t["tag"] for t in r["by"][m]} for m in models}
        common = set.intersection(*sets.values()) if sets else set()
        only = [f"{m}: {', '.join(sorted(sets[m] - set().union(*(sets[o] for o in models if o != m))))}" for m in models
                if sets[m] - set().union(*(sets[o] for o in models if o != m))]
        data.append([f"{r['user']} {r['alias']}", r["ref"] if r["kind"] == "리뷰" else "취향 문장", r["rating"] or "", r["text"], r["product"], r["path"]]
                    + [fmt(r["by"][m]) for m in models] + [", ".join(sorted(common)) or "—", "\n".join(only) or "—"])
    ws = sheet("출처별_비교", head, data, [16, 9, 6, 34, 34, 24] + [24] * len(models) + [18, 30], first=True)
    for k, r in enumerate(rows, start=2):
        if r["rating"] is not None and r["rating"] <= 3:                     # 1~3점 리뷰 — 규칙상 태그가 없어야 하는 행
            for c in ws[k]:
                c.fill = LOW
        elif all({t["tag"] for t in r["by"][m]} == {t["tag"] for t in r["by"][models[0]]} for m in models) and r["by"][models[0]]:
            for c in ws[k][6:6 + len(models)]:
                c.fill = SAME

    summ = []
    for m in models:
        mine = [t for r in rows for t in r["by"][m]]
        pref = sum(len(r["by"][m]) for r in rows if r["ref"] == "pref")
        rev = sum(len(r["by"][m]) for r in rows if r["ref"] != "pref")
        low = sum(len(r["by"][m]) for r in rows if r["rating"] is not None and r["rating"] <= 3)
        empty_hi = sum(1 for r in rows if r["rating"] is not None and r["rating"] >= 4 and not r["by"][m])
        axes: dict[str, int] = {}
        for t in {(x["axis"], x["tag"]) for x in mine}:
            axes[t[0]] = axes.get(t[0], 0) + 1
        nd = sum(len(v) for v in dropped[m].values())
        summ.append([m, len({(r["user"], t["tag"]) for r in rows for t in r["by"][m]}), pref, rev, low, empty_hi,
                     json.dumps(axes, ensure_ascii=False), nd, "같은 프롬프트(2fc950d4)" if m != MANUAL else "대화 중 수작업(고정 프롬프트 없음)"])
    sheet("요약", ["모델", "태그 수(사람별 중복 제거)", "취향 문장에서 나온 태그", "리뷰에서 나온 태그", "1~3점 리뷰에서 나온 태그(0 이어야 함)",
                  "태그 없는 4~5점 리뷰 수", "축 분포(서로 다른 태그)", "검증기가 버린 태그", "조건"], summ, [24, 12, 12, 12, 16, 14, 34, 10, 30])
    sheet("태그_전체", ["사람", "출처", "모델", "축", "태그", "근거 ref(원본)", "비고"], long, [8, 8, 24, 10, 22, 14, 10])
    sheet("설명", ["항목", "내용"], [
        ("무엇", "같은 취향 문장 · 같은 리뷰에서 모델마다 뽑은 태그를 한 행에 나란히 — 출처는 모델이 낸 근거 ref(pref · r1 …)"),
        ("행", "사람 × 출처. 취향 문장이 없는 사람은 리뷰 행만. 리뷰 순서 = 7.6 순서(최신순), 같은 상품 중복 · 카탈로그에 없는 상품은 뺌(프롬프트와 같음)"),
        ("색", "회색 행 = 1~3점 리뷰(규칙상 선호 태그가 없어야 함) · 초록 칸 = 모든 모델이 같은 태그"),
        ("여러 출처", "근거가 \"pref, r1\" 처럼 여럿이면 그 행들에 모두 나온다 — 그래서 행별 태그 수를 더하면 태그 수보다 많을 수 있다"),
        ("claude-opus-5-5", "대화 중 수작업(고정 프롬프트 없음) — 출처는 '태그' 시트 근거 글에서 읽음. 비교 기준은 같은 프롬프트로 다시 뽑은 claude-opus-5-5-prompt"),
        ("원본", "tools/embed_compare/out/태그추출_<모델>_*.json(마지막) · tools/scenario_callback30/scenarios.xlsx"),
    ], [16, 100])
    a.out.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(f"→ {path}  (행 {len(rows)} · 모델 {len(models)})")
    for s in summ:
        print("  " + " · ".join(str(x) for x in s[:8]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
