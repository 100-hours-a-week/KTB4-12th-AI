"""로컬 LLM 서비스 경로 시험 결과표 — 30개 콜백 하네스를 태그 추출 끔 / 모델별로 돌린 결과(JSON · 엑셀)를 나란히.

    cd profiling && uv run --with matplotlib python -m tools.bench.llm_service_report <끔.json> <모델1.json> [<모델2.json> …] [--out 폴더]

각 JSON 옆의 같은 이름 .xlsx '태그' 시트에서 태그를 읽는다(10-12 이전 실행은 JSON 에 태그가 없음). 감점 모드만 표로(D 는 두 모드가 같음).
출력: 화면 표 + <out>/llm_service_<시각>.json · 그래프 2장(docs/assets/report/12-·13-).
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

import numpy as np
import sqlalchemy as sa


def load(path: Path) -> dict:
    runs = json.loads(path.read_text())
    tags: dict[tuple[str, str], str] = {}
    xlsx = path.with_suffix(".xlsx")
    if xlsx.exists():
        from openpyxl import load_workbook

        ws = load_workbook(xlsx, read_only=True)["태그"]
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r and r[0]:
                tags[(r[0], r[1])] = (r[2] or "[]", r[4] or "")
    label = next((r.get("prompt_version") for r in runs if r.get("prompt_version")), None)
    if label is None:
        label = next((v[1] for v in tags.values() if "@" in str(v[1])), "태그 추출 꺼짐")
    return {"path": path, "runs": runs, "tags": tags, "label": str(label).split("@")[-1] if "@" in str(label) else "끔(태그 없음)"}


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("files", nargs="+", type=Path)
    ap.add_argument("--out", type=Path, default=Path("tools/bench/out"))
    a = ap.parse_args()
    from profiling.settings import get_settings

    eng = sa.create_engine(get_settings().DATABASE_URL)
    with eng.connect() as c:
        info = {r[0]: (r[1], r[2]) for r in c.execute(sa.text(
            "select p.id, c.name, p.view_count from ai_catalog.products p "
            "join ai_catalog.categories c on c.id = p.category_id"))}
    views = np.array(sorted(v for _, v in info.values()))
    vpct = lambda v: float(np.searchsorted(views, v, side="left") / (len(views) - 1))
    sets = [load(f) for f in a.files]
    base = sets[0]
    pen = lambda s: {r["user"]: r for r in s["runs"] if r["mode"] == "penalty"}
    bpen = pen(base)
    rows, summary = [], []
    for s in sets:
        P = pen(s)
        srv = []
        for u, r in sorted(P.items()):
            ids = r["callback"]["recommendedProductIds"] if r["callback"] else []
            cnt = Counter(info[i][0] for i in ids)
            top = cnt.most_common(1)[0] if cnt else ("-", 0)
            biased = {x["productId"] for x in r["top30"] if x["bias"] > 0}
            dis = sum(1 for x in r["top30"] if x["penalty"] < 0)
            tg = s["tags"].get((u, "penalty"), ("[]", ""))[0]
            ntag = 0 if tg in ("[]", "") else len(tg.split(" · "))
            ok = all(c["ok"] for c in r["checks"])
            overlap = len(set(ids) & set(bpen[u]["callback"]["recommendedProductIds"])) if bpen.get(u, {}).get("callback") else None
            server = r["timing_ms"]["server"]
            if ntag or s is not base:
                srv.append(server)
            rows.append({"config": s["label"], "user": u, "tags": tg, "n_tags": ntag, "species": len(cnt), "top": f"{top[0]} {top[1]}",
                         "top_n": top[1], "biased": len(biased), "disliked": dis, "overlap_base": overlap,
                         "view_pct": round(float(np.mean([vpct(info[i][1]) for i in ids])), 2) if ids else None,
                         "server_ms": server, "checks_ok": ok, "rule": r.get("rule", "")})
        mine = [x for x in rows if x["config"] == s["label"]]
        tagged = [x for x in mine if x["n_tags"]]
        all_runs = s["runs"]
        checks = [c for r in all_runs for c in r["checks"]]
        sv = [r["timing_ms"]["server"] for r in all_runs]
        summary.append({"config": s["label"], "file": s["path"].name, "checks": f"{sum(c['ok'] for c in checks)}/{len(checks)}",
                        "tagged_users": len(tagged), "tags": sum(x["n_tags"] for x in mine),
                        "species_avg_tagged": round(float(np.mean([x["species"] for x in tagged])), 1) if tagged else None,
                        "species_avg_all": round(float(np.mean([x["species"] for x in mine])), 1),
                        "top_max": max(x["top_n"] for x in mine), "disliked_total": sum(x["disliked"] for x in mine),
                        "overlap_base_avg_tagged": round(float(np.mean([x["overlap_base"] for x in tagged])), 1) if tagged else None,
                        "server_p75": round(float(np.percentile(sv, 75)), 1), "server_p90": round(float(np.percentile(sv, 90)), 1),
                        "server_p95": round(float(np.percentile(sv, 95)), 1), "server_max": round(max(sv), 1)})
    print(f"\n=== 로컬 LLM 서비스 경로 — 기준 {base['label']} · 감점 모드 ===")
    for s in summary:
        print("  " + " · ".join(f"{k} {v}" for k, v in s.items() if k != "file"))
    for cfg in [s["label"] for s in sets]:
        print(f"\n  [{cfg}]")
        for x in (r for r in rows if r["config"] == cfg):
            print(f"   {x['user']} 태그 {x['n_tags']} {x['tags'][:60]:<60} | {x['species']:>2}종 최다 {x['top']:<14} 편향 {x['biased']:>2} 비선호 {x['disliked']} "
                  f"끔과 겹침 {x['overlap_base']} 조회 {x['view_pct']} 서버 {x['server_ms'] / 1000:.1f}s {'✓' if x['checks_ok'] else '✗'}")
    a.out.mkdir(parents=True, exist_ok=True)
    out = a.out / f"llm_service_{datetime.now().astimezone():%Y%m%d_%H%M%S}.json"
    out.write_text(json.dumps({"summary": summary, "rows": rows}, ensure_ascii=False, indent=1))
    try:
        _charts(summary, rows, [s["label"] for s in sets])
    except ImportError:
        print("matplotlib 없음 — 그래프 생략(uv run --with matplotlib)")
    print(f"\n→ {out}")
    return 0


def _charts(summary, rows, labels) -> None:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    plt.rcParams.update({"font.family": "Apple SD Gothic Neo", "font.size": 10, "axes.spines.top": False, "axes.spines.right": False,
                         "axes.unicode_minus": False, "figure.dpi": 150})
    outdir = Path("docs/assets/report")
    col = ["#9AA0A6", "#007D77", "#E8A33D", "#6C5CE7"]
    users = sorted({r["user"] for r in rows if r["user"] not in ("U07", "U08")})
    fig, axs = plt.subplots(2, 1, figsize=(10, 6), sharex=True)
    w = 0.8 / len(labels)
    for k, lab in enumerate(labels):
        x = np.arange(len(users)) + (k - (len(labels) - 1) / 2) * w
        get = {r["user"]: r for r in rows if r["config"] == lab}
        axs[0].bar(x, [get[u]["species"] for u in users], w, color=col[k % 4], label=lab)
        axs[1].bar(x, [get[u]["overlap_base"] for u in users], w, color=col[k % 4])
    axs[0].set_ylabel("30개 안 소분류 종수"); axs[0].legend(frameon=False, fontsize=8, ncol=len(labels), loc="upper left")
    axs[0].set_title("서비스 경로(7.6 → 태그 추출 → D → 7.7) — 감점 모드")
    axs[1].set_ylabel(f"{labels[0]} 와 겹침(/30)"); axs[1].set_xticks(np.arange(len(users)), users); axs[1].set_ylim(0, 32)
    fig.tight_layout(); fig.savefig(outdir / "12-로컬LLM-서비스-섞임.png", bbox_inches="tight"); plt.close(fig)
    fig, ax = plt.subplots(figsize=(7, 3.2))
    ps = ("server_p75", "server_p90", "server_p95")
    pc = ["#B7D7D5", "#5FA8A3", "#007D77"]
    for k, p in enumerate(ps):
        x = np.arange(len(summary)) + (k - 1) * 0.25
        vals = [s[p] / 1000 for s in summary]
        ax.bar(x, vals, 0.25, color=pc[k], label=p.replace("server_", ""))
        for xi, v in zip(x, vals):
            ax.text(xi, v + 0.2, f"{v:.1f}" if v >= 1 else f"{v * 1000:.0f}ms", ha="center", fontsize=7)
    ax.set_xticks(np.arange(len(summary)), [s["config"] for s in summary]); ax.set_ylabel("초")
    ax.set_title("서버 처리 시간(태그 추출 포함, 20건)"); ax.legend(frameon=False, fontsize=8)
    fig.tight_layout(); fig.savefig(outdir / "13-로컬LLM-서비스-시간.png", bbox_inches="tight"); plt.close(fig)
    print("그래프: docs/assets/report/12-로컬LLM-서비스-섞임.png · 13-로컬LLM-서비스-시간.png")


if __name__ == "__main__":
    sys.exit(main())
